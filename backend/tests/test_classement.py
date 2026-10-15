"""Total pondéré et classement du comparatif (15/10/2026) — calculés par le code, données FICTIVES."""

from io import BytesIO

import openpyxl
import pytest

from routers import compare
from services import classement as cl
from services import grille_analyse as ga

GRILLE = {"feuille": "Critères techniques", "ligne_entete": 2, "note_max": 20.0,
          "blocs": [{"nom_cellule": "C1", "col_avis": 3, "col_note": 4},
                    {"nom_cellule": "E1", "col_avis": 5, "col_note": 6}],
          "criteres": [{"ligne": 3, "libelle": "Délais", "documents": "", "points": 60.0},
                       {"ligne": 4, "libelle": "Moyens", "documents": "", "points": 40.0}]}


def _g(nom, n3, n4, **kw):
    return {"nom": nom, "evaluations": {"3": {"avis": "a", "note": n3}, "4": {"avis": "b", "note": n4}}, **kw}


class TestCalcul:
    def test_total_pondere_et_rang(self):
        c = cl.calculer([_g("Alpha", 10, 20), _g("Beta", 20, 10)], [], GRILLE)
        # Alpha : 10/20×60 + 20/20×40 = 70 ; Beta : 60 + 20 = 80
        assert [(l["nom"], l["total"], l["rang"]) for l in c["lignes"]] == [("Beta", 80.0, 1), ("Alpha", 70.0, 2)]
        assert c["max"] == 100.0 and c["pondere"] and c["lignes"][0]["pct"] == 80.0

    def test_ex_aequo_et_note_manquante(self):
        c = cl.calculer([_g("A", 20, 20), _g("B", 20, 20), _g("C", None, 20)], [], GRILLE)
        assert [l["rang"] for l in c["lignes"]] == [1, 1, 3]
        incomplet = c["lignes"][2]
        assert incomplet["nom"] == "C" and incomplet["incomplet"] and incomplet["notes_manquantes"] == 1

    def test_sans_points_somme_des_notes(self):
        grille = {**GRILLE, "criteres": [{**c, "points": None} for c in GRILLE["criteres"]]}
        c = cl.calculer([_g("A", 15, 12)], [], grille)
        assert c["lignes"][0]["total"] == 27.0 and c["max"] == 40.0 and not c["pondere"]

    def test_mode_classique(self):
        groupes = [{"nom": "A", "valeurs": {"Prix (40 pts)": "15/20 — correct", "Délais (60 pts)": "Bon délai"}}]
        c = cl.calculer(groupes, ["Prix (40 pts)", "Délais (60 pts)"], None)
        assert c["lignes"][0]["total"] == 30.0 and c["lignes"][0]["notes_manquantes"] == 1

    def test_rien_a_classer(self):
        assert cl.calculer([{"nom": "A", "valeurs": {"Objet": "texte"}}], ["Objet"], None) is None


class TestRendus:
    def test_markdown_et_resultat(self):
        c = cl.calculer([_g("Alpha", 10, 20), _g("Beta", 20, None)], [], GRILLE)
        md = compare._construire_markdown(["Délais", "Moyens"], [_g("Alpha", 10, 20)], classement=c)
        assert "## 🏆 Classement provisoire" in md and "🥇 1" in md and "⚠️ 1 note(s) manquante(s)" in md

    def test_feuille_excel_formules_vers_la_grille(self, tmp_path):
        wb = openpyxl.Workbook()
        wb.active.title = "Critères techniques"
        blocs = [ga.Bloc(**b) for b in GRILLE["blocs"]]
        groupes = [_g("Alpha", 10, 20), _g("Beta", 20, 10)]
        cl.ajouter_feuille_classement(wb, cl.calculer(groupes, [], GRILLE), GRILLE, groupes, blocs)
        ws = openpyxl.load_workbook(BytesIO(_sauver(wb)))["Classement"]
        assert ws["B2"].value == "Alpha"
        assert ws["C2"].value == "=ROUND(('Critères techniques'!D3*60+'Critères techniques'!D4*40)/20,2)"
        assert ws["C3"].value.startswith("=ROUND(('Critères techniques'!F3*60")
        assert ws["A2"].value == "=RANK(C2,$C$2:$C$3)" and ws["D2"].value == 100

    def test_excel_generique_et_word(self):
        groupes = [{"nom": "A", "valeurs": {"Prix (40 pts)": "15/20 — ok"}}]
        c = cl.calculer(groupes, ["Prix (40 pts)"], None)
        wb = openpyxl.load_workbook(BytesIO(compare._generer_xlsx(None, groupes, ["Prix (40 pts)"], None, classement=c)))
        assert wb["Classement"]["C2"].value == 30.0
        from docx import Document as DocxDocument
        doc = DocxDocument(BytesIO(compare._generer_docx(["Prix (40 pts)"], groupes, None, "T", classement=c)))
        assert any(p.text == "Classement provisoire" for p in doc.paragraphs)


def _sauver(wb):
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_resultat_porte_le_classement(client, db_session):
    from models.job import Job
    j = Job(type="comparatif", statut="completed", parametres={"grille": GRILLE},
            resultat={"colonnes": ["Délais", "Moyens"], "groupes": [_g("Alpha", 10, 20), _g("Beta", 20, 10)]})
    db_session.add(j)
    await db_session.commit()
    r = (await client.get(f"/api/generate/compare/resultat/{j.id}")).json()
    assert r["classement"]["lignes"][0]["nom"] == "Beta" and "Classement provisoire" in r["markdown"]
