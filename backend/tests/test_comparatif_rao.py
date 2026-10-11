"""
Tableau comparatif « analyse d'offres » (11/10/2026) — données 100 % FICTIVES.

1. Une grille RAO aux intitulés différents de MA26001 (« Commentaire », « Notation »…) est reconnue ;
   une simple liste de critères en lignes aussi — au lieu du refus « aucune colonne en ligne 1 ».
2. Fiche points forts / faibles par candidat : lecture tolérante, rendus md / docx / xlsx.
3. Le prompt d'analyse tient dans `num_ctx` et porte les consignes de l'analyste en tête ET en fin.
"""

import uuid
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch

import openpyxl
import pytest

from routers import compare
from services import compare_jobs, fiches_candidats as fc
from services import grille_analyse as ga


# ─── Grilles fictives ────────────────────────────────────────────────────────

def _grille_variante(chemin):
    """Grille avec blocs par candidat, mais d'autres mots que MA26001 et un en-tête en ligne 5."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Feuil1"
    ws["A2"] = "RAPPORT D'ANALYSE — MARCHÉ FICTIF"
    ws["D4"], ws["G4"] = "Entreprise 1", "Entreprise 2"
    ws.merge_cells("D4:F4")
    ws.merge_cells("G4:I4")
    for col, v in zip("ABCDEFGHI", ["Critères techniques", "Pièces", "Pondération", "Commentaire",
                                    "Notation /10", "Note pondérée", "Commentaire", "Notation /10",
                                    "Note pondérée"]):
        ws[f"{col}5"] = v
    for r, (lib, pts) in enumerate([("Méthodologie", 40), ("Moyens humains", 20)], start=6):
        ws[f"A{r}"], ws[f"C{r}"] = lib, pts
        ws[f"F{r}"], ws[f"I{r}"] = f"=E{r}*C{r}/10", f"=H{r}*C{r}/10"
    wb.save(chemin)
    return chemin


def _liste_criteres(chemin):
    """Grille sans bloc candidat : titre fusionné, puis « N° | Critère | Pondération » en ligne 4."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B2"] = "Critères d'attribution"
    ws.merge_cells("B2:D2")
    ws["B4"], ws["C4"], ws["D4"] = "N°", "Critère", "Pondération (points)"
    lignes = [("1", "Valeur technique\nDétail du sous-critère", 60), ("2", "Prix", 40), ("", "TOTAL", 100)]
    for r, (n, lib, pts) in enumerate(lignes, start=5):
        ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"] = n, lib, pts
    wb.save(chemin)
    return chemin


class TestReconnaissanceGrille:
    def test_grille_aux_intitules_differents(self, tmp_path):
        g = ga.detecter_grille(_grille_variante(tmp_path / "g.xlsx"))
        assert g is not None and g.ligne_entete == 5
        assert [(b.nom_cellule, b.col_avis, b.col_note) for b in g.blocs] == [("D4", 4, 5), ("G4", 7, 8)]
        assert [(c.titre, c.points) for c in g.criteres] == [("Méthodologie", 40.0), ("Moyens humains", 20.0)]

    def test_liste_de_criteres_en_lignes(self, tmp_path):
        chemin = _liste_criteres(tmp_path / "l.xlsx")
        assert ga.detecter_grille(chemin) is None          # pas de bloc candidat
        liste = ga.lire_liste_criteres(chemin)
        assert [c.colonne for c in liste] == ["Valeur technique (60 pts)", "Prix (40 pts)"]   # TOTAL écarté

    def test_ancien_classeur_en_colonnes_inchange(self, tmp_path):
        wb = openpyxl.Workbook()
        wb.active.append(["Société", "Prix", "Délai"])
        wb.save(tmp_path / "t.xlsx")
        assert ga.lire_liste_criteres(tmp_path / "t.xlsx") == []
        apercu = ga.apercu_classeur(tmp_path / "t.xlsx")
        assert apercu["type"] == "colonnes" and [c["nom"] for c in apercu["criteres"]] == ["Société", "Prix", "Délai"]

    @pytest.mark.parametrize("fabrique, attendu", [(_grille_variante, "grille"), (_liste_criteres, "liste")])
    def test_apercu(self, tmp_path, fabrique, attendu):
        assert ga.apercu_classeur(fabrique(tmp_path / "x.xlsx"))["type"] == attendu

    def test_apercu_classeur_vide(self, tmp_path):
        openpyxl.Workbook().save(tmp_path / "v.xlsx")
        apercu = ga.apercu_classeur(tmp_path / "v.xlsx")
        assert apercu["type"] == "inconnu" and "Je les saisis" in apercu["message"]

    def test_grille_remplie_garde_formules_et_ajoute_les_fiches(self, tmp_path):
        chemin = _grille_variante(tmp_path / "g.xlsx")
        g = ga.detecter_grille(chemin)
        groupes = [{"nom": "Alpha", "evaluations": {"6": {"avis": "Méthode claire.", "note": 8}},
                    "fiche": {"resume": "Offre solide.", "points_forts": ["Méthode"], "points_faibles": [],
                              "vigilance": []}},
                   {"nom": "Beta", "evaluations": {}, "fiche": None}]
        wb = openpyxl.load_workbook(BytesIO(ga.remplir_grille(chemin, g, groupes)))
        ws = wb["Feuil1"]
        assert ws["D6"].value == "Méthode claire." and ws["E6"].value == 8
        assert ws["F6"].value == "=E6*C6/10"                      # formule intacte
        fiches = wb[fc.TITRE_FEUILLE]
        assert fiches["B2"].value == "Alpha" and "Méthode" in fiches["B4"].value
        assert "non produite" in fiches["C3"].value


# ─── Lancement et modèle : liste → critères, plus de refus muet ──────────────

GROUPES_FICTIFS = [{"nom": "A", "document_ids": [str(uuid.uuid4())]},
                   {"nom": "B", "document_ids": [str(uuid.uuid4())]}]


async def _modele(db_session, chemin):
    from models.template import Template
    t = Template(nom="Grille", type="xlsx", chemin_fichier=str(chemin))
    db_session.add(t)
    await db_session.commit()
    return t


class TestLancement:
    @pytest.mark.asyncio
    async def test_liste_de_criteres_acceptee(self, client, db_session, tmp_path):
        t = await _modele(db_session, _liste_criteres(tmp_path / "l.xlsx"))
        r = await client.post("/api/generate/compare", json={"template_id": str(t.id), "groupes": GROUPES_FICTIFS})
        assert r.status_code == 202, r.text
        assert r.json()["colonnes"] == ["Valeur technique (60 pts)", "Prix (40 pts)"]

    @pytest.mark.asyncio
    async def test_classeur_illisible_message_explicite(self, client, db_session, tmp_path):
        openpyxl.Workbook().save(tmp_path / "v.xlsx")
        t = await _modele(db_session, tmp_path / "v.xlsx")
        r = await client.post("/api/generate/compare", json={"template_id": str(t.id), "groupes": GROUPES_FICTIFS})
        assert r.status_code == 400 and "Je les saisis" in r.json()["detail"]

    @pytest.mark.asyncio
    async def test_fichier_et_apercu_du_modele(self, client, db_session, tmp_path):
        chemin = _liste_criteres(tmp_path / "l.xlsx")
        t = await _modele(db_session, chemin)
        f = await client.get(f"/api/templates/{t.id}/fichier")
        assert f.status_code == 200 and f.content == chemin.read_bytes()
        a = await client.get(f"/api/templates/{t.id}/apercu")
        assert a.status_code == 200 and a.json()["type"] == "liste"


# ─── Fiches candidats ────────────────────────────────────────────────────────

class TestFiches:
    def test_lecture_tolerante(self):
        reponse = ('```json\n{"resume": "Offre\ncomplète.", "points_forts": ["- **Délais** tenus", ""],'
                   ' "points_faibles": "Prix élevé\\nPeu de références", "vigilance": []}\n```')
        fiche = fc.lire_fiche(reponse)
        assert fiche["resume"] == "Offre\ncomplète."
        assert fiche["points_forts"] == ["Délais tenus"]
        assert fiche["points_faibles"] == ["Prix élevé", "Peu de références"]

    @pytest.mark.parametrize("reponse", ["rien", "{}", '{"resume": "", "points_forts": []}', "{pas du json"])
    def test_fiche_inexploitable(self, reponse):
        assert fc.lire_fiche(reponse) is None

    def test_rien_a_resumer_pas_d_appel(self):
        assert fc.construire_prompt("A", {"Prix": "N/A", "Délai": "⚠ Échec de l'IA"}, None) is None

    @pytest.mark.asyncio
    async def test_appel_ia_en_erreur_donne_none(self):
        o = MagicMock()
        o.generate = AsyncMock(side_effect=RuntimeError("Ollama injoignable"))
        assert await fc.rediger_fiche("A", {"Prix": "10 €"}, None, "m", o) is None

    def test_rendus(self):
        groupes = [{"nom": "A", "valeurs": {"Prix": "10"},
                    "fiche": {"resume": "Bien.", "points_forts": ["Prix bas"], "points_faibles": ["Délai"],
                              "vigilance": []}},
                   {"nom": "B", "valeurs": {"Prix": "12"}, "fiche": None}]
        md = compare._construire_markdown(["Prix"], groupes)
        assert "## Fiches candidats" in md and "- Prix bas" in md and "non produite" in md
        assert "Fiches candidats" not in compare._construire_markdown(["Prix"], [{"nom": "A", "valeurs": {}}])
        wb = openpyxl.load_workbook(BytesIO(compare._generer_xlsx(None, groupes, ["Prix"], None)))
        assert fc.TITRE_FEUILLE in wb.sheetnames
        from docx import Document as DocxDocument
        doc = DocxDocument(BytesIO(compare._generer_docx(["Prix"], groupes, None, "T")))
        assert any(p.text == "Prix bas" for p in doc.paragraphs)


# ─── Prompt d'analyse : budget et consignes ──────────────────────────────────

class TestPromptAnalyse:
    @pytest.mark.asyncio
    async def test_tient_dans_num_ctx_et_consignes_en_tete_et_en_fin(self, monkeypatch):
        monkeypatch.setattr(compare.settings, "ollama_num_ctx", 16384)
        colonnes = [f"Critère {i}" for i in range(25)]
        consignes = "Tu es ingénieur territorial. " * 60          # ≈ 1 700 caractères
        docs = [SimpleNamespace(nom=f"offre{i}.pdf", texte_extrait="mot " * 50_000) for i in range(3)]
        o = MagicMock()
        o.generate = AsyncMock(return_value="{}")
        await compare._analyser_groupe("Alpha", docs, colonnes, consignes, "m", o)
        prompt = o.generate.await_args.args[0]
        tokens = len(prompt) / compare.CHARS_PAR_TOKEN + compare.TOKENS_REPONSE_PAR_COLONNE * len(colonnes)
        assert tokens <= 16384
        assert prompt.startswith("=== CONSIGNES DE L'ANALYSTE ===")
        assert "ingénieur territorial" in prompt.split("=== FORMAT")[0]
        assert prompt.rstrip().endswith("la réponse reste uniquement cet objet JSON.")

    def test_part_lue(self):
        docs = [SimpleNamespace(nom="a.pdf", texte_extrait="x" * 100_000)]
        assert compare._part_lue(docs, 30_000) < 100
        assert compare._part_lue([SimpleNamespace(nom="b.pdf", texte_extrait="court")], 30_000) == 100


# ─── Handler : fiche rédigée après chaque candidat ───────────────────────────

class TestHandlerFiches:
    @pytest.mark.asyncio
    async def test_fiche_par_candidat_et_part_lue(self):
        ctx = MagicMock()
        ctx.job_id = str(uuid.uuid4())
        ctx.parametres = {"groupes": [{"nom": "A", "document_ids": []}, {"nom": "B", "document_ids": []}],
                          "colonnes": ["Prix"], "model": "m", "synthese": False}
        ctx.cancelled = False
        ctx.report = AsyncMock()
        analyses = AsyncMock(side_effect=[({"Prix": "10"}, True), ({"Prix": compare.VALEUR_ECHEC_IA}, False)])
        fiche = {"resume": "r", "points_forts": ["p"], "points_faibles": [], "vigilance": []}
        rediger = AsyncMock(return_value=fiche)
        with patch.object(compare, "_analyser_groupe", analyses), \
             patch.object(compare, "_charger_documents", AsyncMock(return_value=[])), \
             patch.object(fc, "rediger_fiche", rediger), \
             patch.object(compare_jobs, "_persister", AsyncMock()):
            etat = await compare_jobs.handler_comparatif(ctx)
        assert etat["groupes"][0]["fiche"] == fiche and etat["groupes"][0]["lu_pct"] == 100
        assert etat["groupes"][1]["fiche"] is None          # candidat en échec : pas de fiche inventée
        assert rediger.await_count == 1


# ─── Disposition MA26041 (structure seule, 11/10/2026) ──────────────────────

def _grille_ma26041(chemin):
    """5 candidats « centrés sur plusieurs colonnes » (pas de fusion), 2 paires Avis / Note sur 20
    par candidat, lignes « TOTAL note provisoire » / « TOTAL note finale recalculée »."""
    from openpyxl.styles import Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Critères techniques"
    ws["A1"] = "Candidat"
    for i, col in enumerate("CGKOS", start=1):
        ws[f"{col}1"] = f"candidat #{i:02d}"
        ws[f"{col}1"].alignment = Alignment(horizontal="centerContinuous")
    ws["A2"], ws["B2"] = "Sous-critère technique", "Points par sous critère"
    for c in range(3, 23, 2):
        ws.cell(row=2, column=c, value="Avis")
        ws.cell(row=2, column=c + 1, value="Note sur 20")
        ws.cell(row=6, column=c, value="TOTAL note provisoire")
        ws.cell(row=7, column=c, value="TOTAL note finale recalculée")
    for r, lib in enumerate(["Critère A", "Critère B", "Critère C"], start=3):
        ws[f"A{r}"], ws[f"B{r}"] = lib, 20
    wb.save(chemin)
    return chemin


class TestGrilleMa26041:
    def test_un_bloc_par_candidat_et_notes_sur_20(self, tmp_path):
        g = ga.detecter_grille(_grille_ma26041(tmp_path / "g.xlsx"))
        assert [(b.nom_cellule, b.col_avis, b.col_note) for b in g.blocs] == [
            ("C1", 3, 4), ("G1", 7, 8), ("K1", 11, 12), ("O1", 15, 16), ("S1", 19, 20)]
        assert g.note_max == 20.0
        assert ga.Grille.from_dict(g.to_dict()) == g
        assert ga.Grille.from_dict({**g.to_dict(), "note_max": None}).note_max == 10.0   # analyses d'avant

    def test_remplissage_noms_notes_et_total_provisoire(self, tmp_path):
        chemin = _grille_ma26041(tmp_path / "g.xlsx")
        g = ga.detecter_grille(chemin)
        groupes = [{"nom": "Alpha", "evaluations": {"3": {"avis": "Délais garantis.", "note": 17},
                                                    "4": {"avis": "Équipe dédiée.", "note": 14}}},
                   {"nom": "Beta", "evaluations": {"3": {"avis": "Flou.", "note": 8}}}]
        ws = openpyxl.load_workbook(BytesIO(ga.remplir_grille(chemin, g, groupes)))["Critères techniques"]
        assert ws["C1"].value == "Alpha" and ws["G1"].value == "Beta"      # « candidat #0x » renommés
        assert ws["K1"].value == "candidat #03"                              # bloc sans groupe : intact
        assert ws["C3"].value == "Délais garantis." and ws["D3"].value == 17
        assert ws["E3"].value is None and ws["F3"].value is None            # 2e paire : à la commission
        assert ws["D6"].value == "=SUM(D3:D5)" and ws["D7"].value is None    # total provisoire seul
        assert ws["H6"].value == "=SUM(H3:H5)"

    @pytest.mark.asyncio
    async def test_note_bornee_a_l_echelle(self):
        o = MagicMock()
        o.generate = AsyncMock(return_value='{"note": 25, "avis": "Très complet."}')
        critere = ga.Critere(ligne=3, libelle="Critère A", documents="", points=20)
        avis, note, ok = await ga.evaluer_critere("A", critere, [], None, "m", o, contexte="x", note_max=20)
        assert ok and note == 20.0
        assert "de 0 à 20" in o.generate.await_args.args[0]
        assert ga.valeur_lisible("Bien.", 15, 20) == "15/20 — Bien."


# ─── Documents du marché (CCTP, règlement) confrontés aux offres ─────────────

class TestDocumentsMarche:
    @pytest.mark.asyncio
    async def test_exigences_dans_le_prompt_du_critere(self):
        o = MagicMock()
        o.generate = AsyncMock(return_value='{"note": 12, "avis": "Répond en partie."}')
        critere = ga.Critere(ligne=3, libelle="Délais d'urgence", documents="", points=20)
        await ga.evaluer_critere("A", critere, [], None, "m", o, contexte="OFFRE", note_max=20,
                                 exigences="Intervention sous 4 h exigée (CCTP art. 5).")
        prompt = o.generate.await_args.args[0]
        assert "Exigences du marché sur ce critère" in prompt and "sous 4 h" in prompt
        assert prompt.index("sous 4 h") < prompt.index("OFFRE")                   # exigences avant l'offre
        assert "y répond en partie, ou n'en dit rien" in prompt

    @pytest.mark.asyncio
    async def test_sans_document_du_marche_rien_ne_change(self):
        o = MagicMock()
        o.generate = AsyncMock(return_value='{"note": 12, "avis": "x"}')
        critere = ga.Critere(ligne=3, libelle="Délais", documents="", points=20)
        await ga.evaluer_critere("A", critere, [], None, "m", o, contexte="OFFRE")
        assert "Exigences du marché" not in o.generate.await_args.args[0]
        assert await ga.exigences_marche(critere, []) is None

    @pytest.mark.asyncio
    async def test_exigences_repli_sans_vecteurs(self):
        critere = ga.Critere(ligne=3, libelle="Délais", documents="", points=20)
        cctp = [SimpleNamespace(id=uuid.uuid4(), nom="CCTP.pdf", texte_extrait="Article 5 : intervention sous 4 h.")]
        with patch.object(ga, "passages_pertinents", AsyncMock(return_value=None)):
            assert "intervention sous 4 h" in await ga.exigences_marche(critere, cctp)

    @pytest.mark.asyncio
    async def test_analyse_classique_recoit_les_exigences(self):
        o = MagicMock()
        o.generate = AsyncMock(return_value='{"Prix": "10"}')
        await compare._analyser_groupe("A", [], ["Prix"], None, "m", o, exigences="Prix ferme exigé.")
        prompt = o.generate.await_args.args[0]
        assert prompt.startswith("=== EXIGENCES DU MARCHÉ") and "Prix ferme exigé." in prompt

    @pytest.mark.asyncio
    async def test_handler_grille_exigences_une_fois_par_critere(self, tmp_path):
        g = ga.detecter_grille(_grille_ma26041(tmp_path / "g.xlsx"))
        ctx = MagicMock()
        ctx.job_id = str(uuid.uuid4())
        ctx.parametres = {"groupes": [{"nom": "A", "document_ids": []}, {"nom": "B", "document_ids": []}],
                          "colonnes": [c.titre for c in g.criteres], "model": "m", "synthese": False,
                          "fiches": False, "grille": g.to_dict(), "documents_marche": [str(uuid.uuid4())]}
        ctx.cancelled = False
        ctx.report = AsyncMock()
        cctp = [SimpleNamespace(id=uuid.uuid4(), nom="CCTP.pdf", texte_extrait="Exigence.")]
        charger = AsyncMock(side_effect=lambda ids: cctp if ids else [])
        exigences = AsyncMock(return_value="EXIGENCE")
        evaluer = AsyncMock(return_value=("Avis.", 15.0, True))
        with patch.object(compare, "_charger_documents", charger), \
             patch.object(ga, "exigences_marche", exigences), \
             patch.object(ga, "passages_pertinents", AsyncMock(return_value="OFFRE")), \
             patch.object(ga, "evaluer_critere", evaluer), \
             patch.object(compare_jobs, "_persister", AsyncMock()):
            etat = await compare_jobs.handler_comparatif(ctx)
        assert exigences.await_count == len(g.criteres)                     # pas une fois par candidat
        assert all(c.kwargs["exigences"] == "EXIGENCE" for c in evaluer.await_args_list)
        assert any("Documents du marché" in (e.get("message") or "") for e in etat["events"])
