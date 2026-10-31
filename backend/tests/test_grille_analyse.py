"""Grille d'analyse des offres : détection, évaluation par critère, remplissage du classeur fourni."""
from io import BytesIO
from types import SimpleNamespace

import openpyxl
import pytest

from services import grille_analyse as ga


def _grille_rao(chemin):
    """Reproduit la disposition de MA26001_RAO_Lot1.xlsx (01/10/2026)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Critères techniques"
    ws["B1"] = "CRITERES D'ANALYSE DES OFFRES"
    ws["E2"], ws["H2"] = "{{NOM_SOCIETE}}", "{{NOM_SOCIETE 2 }}"
    ws.merge_cells("E2:G2")
    ws.merge_cells("H2:J2")
    for col, v in zip("BCDEFGHIJ", ["Libellé", "Documents analysés", "Points par sous critère",
                                    "Avis", "Note sur 10", "Note finale", "Avis", "Note sur 10", "Note finale"]):
        ws[f"{col}3"] = v
    lignes = [
        ("2.1-Délais, phasage et organisation de chantier\no Les délais…", "• Acte d’engagement,\n• Planning", 10),
        ("2.2- Qualité technique du projet\no La qualité…", "• Mémoire technique 6.1", 15),
    ]
    for r, (lib, docs, pts) in enumerate(lignes, start=4):
        ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"] = lib, docs, pts
        ws[f"G{r}"], ws[f"J{r}"] = f"=F{r}*D{r}/10", f"=D{r}*I{r}/10"
    wb.save(chemin)
    return chemin


def test_detecte_la_grille_rao(tmp_path):
    g = ga.detecter_grille(_grille_rao(tmp_path / "rao.xlsx"))
    assert g is not None
    assert g.ligne_entete == 3
    assert [(b.nom_cellule, b.col_avis, b.col_note) for b in g.blocs] == [("E2", 5, 6), ("H2", 8, 9)]
    assert [c.titre for c in g.criteres] == ["2.1-Délais, phasage et organisation de chantier",
                                             "2.2- Qualité technique du projet"]
    assert g.criteres[0].points == 10.0
    assert "Planning" in g.criteres[0].documents
    assert ga.Grille.from_dict(g.to_dict()) == g


def test_classeur_ordinaire_pas_une_grille(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(["Société", "Prix", "Délai"])
    wb.save(tmp_path / "t.xlsx")
    assert ga.detecter_grille(tmp_path / "t.xlsx") is None


def test_remplir_conserve_formules_et_remplace_noms(tmp_path):
    chemin = _grille_rao(tmp_path / "rao.xlsx")
    g = ga.detecter_grille(chemin)
    groupes = [
        {"nom": "CES", "evaluations": {"4": {"avis": "Planning détaillé.", "note": 7.5},
                                       "5": {"avis": ga.AVIS_ECHEC_IA, "note": None}}},
        {"nom": "OTV", "evaluations": {"4": {"avis": "Délais flous.", "note": 5}}},
    ]
    ws = openpyxl.load_workbook(BytesIO(ga.remplir_grille(chemin, g, groupes)))["Critères techniques"]
    assert (ws["E2"].value, ws["H2"].value) == ("CES", "OTV")
    assert (ws["E4"].value, ws["F4"].value) == ("Planning détaillé.", 7.5)
    assert (ws["H4"].value, ws["I4"].value) == ("Délais flous.", 5)
    assert ws["E5"].value == ga.AVIS_ECHEC_IA and ws["F5"].value is None   # pas de note inventée
    assert ws["G4"].value == "=F4*D4/10" and ws["J4"].value == "=D4*I4/10"
    assert ws["B4"].value.startswith("2.1")


def test_nom_deja_saisi_conserve(tmp_path):
    chemin = _grille_rao(tmp_path / "rao.xlsx")
    wb = openpyxl.load_workbook(chemin)
    wb.active["E2"] = "Entreprise Martin"
    wb.save(chemin)
    g = ga.detecter_grille(chemin)
    ws = openpyxl.load_workbook(BytesIO(ga.remplir_grille(chemin, g, [{"nom": "CES"}]))).active
    assert ws["E2"].value == "Entreprise Martin"


def test_pieces_citees_en_tete():
    docs = [SimpleNamespace(nom="7-Attestation de Visite.pdf"),
            SimpleNamespace(nom="6_2-Planning AO Lot_01.pdf"),
            SimpleNamespace(nom="1-Lot1 Acte engagement.pdf")]
    ordre = [d.nom for d in ga._ordonner_documents(docs, "• Acte d’engagement,\n• Planning")]
    assert ordre[-1] == "7-Attestation de Visite.pdf"


class _OllamaFactice:
    def __init__(self, reponse):
        self.reponse, self.prompts = reponse, []

    async def generate(self, prompt, model=None, num_predict=None):
        self.prompts.append(prompt)
        self.num_predict = num_predict
        if isinstance(self.reponse, Exception):
            raise self.reponse
        return self.reponse


CRITERE = ga.Critere(ligne=4, libelle="2.1-Délais\no Les délais", documents="• Planning", points=10)
DOCS = [SimpleNamespace(nom="Planning.pdf", texte_extrait="Durée 12 semaines.")]


@pytest.mark.parametrize("reponse, attendu", [
    ('Voici : {"avis": "Bon planning.", "note": "7,5"}', ("Bon planning.", 7.5, True)),
    ('{"avis": "Excellent", "note": 14}', ("Excellent", 10.0, True)),
    ("Je ne sais pas.", (ga.AVIS_ECHEC_IA, None, False)),
    (RuntimeError("Ollama injoignable"), (ga.AVIS_ECHEC_IA, None, False)),
])
async def test_evaluer_critere(reponse, attendu):
    ollama = _OllamaFactice(reponse)
    assert await ga.evaluer_critere("CES", CRITERE, DOCS, None, "m", ollama) == attendu


async def test_consigne_apres_les_documents():
    ollama = _OllamaFactice('{"avis": "ok", "note": 5}')
    await ga.evaluer_critere("CES", CRITERE, DOCS, "Sois sévère", "m", ollama)
    prompt = ollama.prompts[0]
    assert prompt.index("Durée 12 semaines.") < prompt.index("=== CONSIGNE ===") < prompt.index("Sois sévère")


def test_valeur_lisible():
    assert ga.valeur_lisible("Bon.", 7.5) == "7.5/10 — Bon."
    assert ga.valeur_lisible(ga.AVIS_ECHEC_IA, None) == ga.AVIS_ECHEC_IA


async def test_passages_repli_sans_vecteurs(monkeypatch):
    """Embedding indisponible → None : l'évaluation retombe sur le début des pièces citées."""
    import routers.search as rs

    async def _aucun(q):
        return None
    monkeypatch.setattr(rs, "_embed_query", _aucun)
    assert await ga.passages_pertinents(CRITERE, [SimpleNamespace(id="x", nom="a.pdf")]) is None


async def test_contexte_fourni_prioritaire():
    ollama = _OllamaFactice('{"avis": "ok", "note": 5}')
    await ga.evaluer_critere("CES", CRITERE, DOCS, None, "m", ollama, contexte="PASSAGE CHOISI")
    assert "PASSAGE CHOISI" in ollama.prompts[0] and "Durée 12 semaines." not in ollama.prompts[0]


# Forme EXACTE des réponses de ministral-3 du 01/10/2026 : retours à la ligne bruts dans l'avis.
REPONSE_REELLE = '```json\n{\n  "avis": "\n  **Points forts :**\n  - **Cohérence des délais** : planning détaillé (6_2-Planning AO Lot_01.pdf).\n\n\n  **Points faibles :**\n  - Dossier plan absent.",\n  "note": 7.5\n}\n```'


def test_reponse_avec_retours_a_la_ligne_bruts():
    avis, note = ga._lire_reponse(REPONSE_REELLE)
    assert note == 7.5
    assert avis.startswith("Points forts :") and "Dossier plan absent." in avis
    assert "**" not in avis and "\n\n\n" not in avis


def test_reponse_guillemets_non_echappes():
    avis, note = ga._lire_reponse('{"avis": "Le mémoire dit "phasage" en 3 temps.", "note": "6"}')
    assert note == 6.0 and "phasage" in avis


def test_reponse_tronquee_avant_la_note():
    avis, note = ga._lire_reponse('```json\n{"avis": "Planning précis, moyens cohérents')
    assert note is None and avis == "Planning précis, moyens cohérents"


def test_from_dict_ignore_la_copie_du_classeur(tmp_path):
    g = ga.detecter_grille(_grille_rao(tmp_path / "rao.xlsx"))
    d = {**g.to_dict(), "fichier": "/app/storage/templates/_analyses/x.xlsx"}
    assert ga.Grille.from_dict(d) == g


async def test_reponse_plafonnee_et_prompt_borne():
    """Le prompt laisse de la place à l'avis, et la génération est plafonnée (incident du 01/10)."""
    ollama = _OllamaFactice('{"avis": "ok", "note": 5}')
    gros = [SimpleNamespace(nom=f"m{i}.pdf", texte_extrait="x" * 50_000) for i in range(4)]
    await ga.evaluer_critere("CES", CRITERE, gros, None, "m", ollama)
    assert ollama.num_predict == ga.TOKENS_AVIS_MAX
    assert len(ollama.prompts[0]) < ga.BUDGET_CRITERE + 3_000


def test_note_en_premier_avis_coupe():
    """Réponse plafonnée : la note, demandée en premier, survit à la coupure de l'avis."""
    avis, note = ga._lire_reponse('```json\n{"note": 6.5, "avis": "Continuité de service bien traitée, mais')
    assert note == 6.5 and avis == "Continuité de service bien traitée, mais"


async def test_consigne_note_en_premier():
    ollama = _OllamaFactice('{"note": 5, "avis": "ok"}')
    await ga.evaluer_critere("CES", CRITERE, DOCS, None, "m", ollama)
    assert '{"note": 7, "avis": "…"}' in ollama.prompts[0]


# ─── Barème lu dans la grille (« 0 = absent… ») — ROADMAP, comparatif step 2 ───────────────────


def _avec_bareme_en_ligne(chemin):
    """Grille RAO + feuille « Notation » : une cellule par niveau, « N = libellé »."""
    _grille_rao(chemin)
    wb = openpyxl.load_workbook(chemin)
    ws = wb.create_sheet("Notation")
    ws["A1"] = "Barème de notation"
    for r, v in enumerate(["0 = Absent ou hors sujet", "4 : Insuffisant", "7 = Satisfaisant",
                           "10 = Très satisfaisant"], start=2):
        ws[f"A{r}"] = v
    wb.save(chemin)
    return chemin


def test_bareme_en_ligne_lu(tmp_path):
    g = ga.detecter_grille(_avec_bareme_en_ligne(tmp_path / "rao.xlsx"))
    assert g.bareme == [{"note": 0.0, "libelle": "Absent ou hors sujet"},
                        {"note": 4.0, "libelle": "Insuffisant"},
                        {"note": 7.0, "libelle": "Satisfaisant"},
                        {"note": 10.0, "libelle": "Très satisfaisant"}]
    assert ga.Grille.from_dict(g.to_dict()) == g


def test_bareme_en_tableau_lu(tmp_path):
    chemin = _grille_rao(tmp_path / "rao.xlsx")
    wb = openpyxl.load_workbook(chemin)
    ws = wb["Critères techniques"]
    ws["L3"] = "Barème"
    for r, (n, lib) in enumerate([(0, "Non traité"), (5, "Partiel"), (10, "Complet")], start=4):
        ws[f"L{r}"], ws[f"M{r}"] = n, lib
    wb.save(chemin)
    g = ga.detecter_grille(chemin)
    assert [(b["note"], b["libelle"]) for b in g.bareme] == [(0.0, "Non traité"), (5.0, "Partiel"),
                                                             (10.0, "Complet")]


def test_sans_bareme(tmp_path):
    g = ga.detecter_grille(_grille_rao(tmp_path / "rao.xlsx"))
    assert g.bareme == []
    # Analyse lancée avant le barème : la clé manque, la grille se relit quand même.
    d = g.to_dict()
    d.pop("bareme")
    assert ga.Grille.from_dict(d).bareme == []


def test_bareme_du_libelle_de_critere_ignore(tmp_path):
    """Un « 0 = … » DANS le libellé d'un critère est déjà envoyé à l'IA avec le critère."""
    chemin = _grille_rao(tmp_path / "rao.xlsx")
    wb = openpyxl.load_workbook(chemin)
    ws = wb["Critères techniques"]
    ws["B4"] = "2.1-Délais\n0 = aucun planning\n10 = planning détaillé"
    wb.save(chemin)
    assert ga.detecter_grille(chemin).bareme == []


def test_echelle_deduite_du_bareme(tmp_path):
    """En-tête « Note » sans « sur N » : le barème (0 à 20) donne l'échelle."""
    chemin = _avec_bareme_en_ligne(tmp_path / "rao.xlsx")
    wb = openpyxl.load_workbook(chemin)
    wb["Critères techniques"]["F3"] = "Note"
    wb["Critères techniques"]["I3"] = "Note"
    wb["Notation"]["A6"] = "20 = Excellent"
    wb.save(chemin)
    g = ga.detecter_grille(chemin)
    assert g.note_max == 20.0
    assert g.bareme[-1] == {"note": 20.0, "libelle": "Excellent"}


def test_bareme_hors_echelle_ecarte(tmp_path):
    """« Note sur 10 » est explicite : un niveau à 20 contredit la grille, il n'est pas repris."""
    chemin = _avec_bareme_en_ligne(tmp_path / "rao.xlsx")
    wb = openpyxl.load_workbook(chemin)
    wb["Notation"]["A6"] = "20 = Excellent"
    wb.save(chemin)
    g = ga.detecter_grille(chemin)
    assert g.note_max == 10.0
    assert all(b["note"] <= 10 for b in g.bareme)


async def test_bareme_dans_le_prompt():
    ollama = _OllamaFactice('{"note": 7, "avis": "ok"}')
    bareme = [{"note": 0.0, "libelle": "Absent"}, {"note": 10.0, "libelle": "Complet"}]
    await ga.evaluer_critere("CES", CRITERE, DOCS, None, "m", ollama, bareme=bareme)
    prompt = ollama.prompts[0]
    assert "0 = Absent" in prompt and "10 = Complet" in prompt
    # Rappelé dans la consigne, après les documents (Ollama coupe le DÉBUT en cas de débordement).
    assert prompt.index("Durée 12 semaines.") < prompt.rindex("barème")


async def test_prompt_sans_bareme_inchange():
    ollama = _OllamaFactice('{"note": 7, "avis": "ok"}')
    await ga.evaluer_critere("CES", CRITERE, DOCS, None, "m", ollama)
    assert "barème" not in ollama.prompts[0].lower()


def test_apercu_annonce_le_bareme(tmp_path):
    a = ga.apercu_classeur(_avec_bareme_en_ligne(tmp_path / "rao.xlsx"))
    assert a["bareme"][0] == {"note": 0.0, "libelle": "Absent ou hors sujet"}
    assert "barème" in a["message"]


# ─── Correction de la grille détectée, avant lancement — ROADMAP, comparatif step 2 ────────────


def test_correction_renomme_repondere_retire(tmp_path):
    g = ga.detecter_grille(_avec_bareme_en_ligne(tmp_path / "rao.xlsx"))
    c = ga.appliquer_correction(g, {
        "criteres": [{"ligne": 4, "nom": "Délais", "points": 30, "actif": True},
                     {"ligne": 5, "nom": "2.2- Qualité technique du projet", "points": 15, "actif": False}],
    })
    assert [cr.ligne for cr in c.criteres] == [4]
    assert c.criteres[0].titre == "Délais"
    assert "Les délais" in c.criteres[0].libelle       # le détail du critère reste envoyé à l'IA
    assert c.criteres[0].points == 30.0
    assert c.bareme == g.bareme                         # barème non fourni : inchangé
    assert g.criteres[0].titre.startswith("2.1")        # la grille d'origine n'est pas modifiée


def test_correction_bareme_et_echelle(tmp_path):
    g = ga.detecter_grille(_grille_rao(tmp_path / "rao.xlsx"))
    c = ga.appliquer_correction(g, {"note_max": 20, "bareme": [
        {"note": 20, "libelle": "Excellent"}, {"note": 0, "libelle": " Absent "},
        {"note": 25, "libelle": "Hors échelle"}, {"note": 10, "libelle": ""}]})
    assert c.note_max == 20.0
    assert c.bareme == [{"note": 0.0, "libelle": "Absent"}, {"note": 20.0, "libelle": "Excellent"}]


def test_correction_bareme_vide_le_retire(tmp_path):
    g = ga.detecter_grille(_avec_bareme_en_ligne(tmp_path / "rao.xlsx"))
    assert ga.appliquer_correction(g, {"bareme": []}).bareme == []


def test_correction_ligne_inconnue_ignoree_et_aucun_critere_refuse(tmp_path):
    g = ga.detecter_grille(_grille_rao(tmp_path / "rao.xlsx"))
    c = ga.appliquer_correction(g, {"criteres": [{"ligne": 99, "nom": "Inventé", "points": 5}]})
    assert [cr.ligne for cr in c.criteres] == [4, 5]
    with pytest.raises(ValueError):
        ga.appliquer_correction(g, {"criteres": [{"ligne": 4, "actif": False}, {"ligne": 5, "actif": False}]})


def test_apercu_donne_ligne_et_echelle(tmp_path):
    a = ga.apercu_classeur(_grille_rao(tmp_path / "rao.xlsx"))
    assert [c["ligne"] for c in a["criteres"]] == [4, 5]
    assert a["note_max"] == 10.0
