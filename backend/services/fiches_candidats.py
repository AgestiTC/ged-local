"""
Fiches candidats du tableau comparatif — points forts, points faibles, vigilance.
================================================================================
Le tableau dit « quoi » critère par critère ; la synthèse des écarts compare les candidats entre
eux. Il manquait, pour chaque candidat, un RÉSUMÉ lisible de son offre (demande du 11/10/2026,
analyse d'offres d'un marché public).

La fiche est rédigée à partir de l'ÉVALUATION déjà faite (valeurs ou avis par critère), pas en
relisant les documents : l'appel est court, tient toujours dans le contexte, et la fiche ne peut
pas contredire le tableau. Une fiche que l'IA n'a pas su produire vaut None, et s'affiche
« non produite » — jamais une fiche vide présentée comme pleine.

Ce module : l'appel IA (`rediger_fiche`), la lecture tolérante de sa réponse (`lire_fiche`) et
les rendus (Markdown, Word, feuille Excel).
"""
from __future__ import annotations

import json
import re

from logger import get_logger

log = get_logger(__name__)

# Résumé + 3 listes de quelques puces. 900 (11/10/2026) laissait des JSON coupés avant la fin :
# 2 fiches sur 7 « inexploitables » alors que l'évaluation était bonne.
TOKENS_FICHE_MAX = 1_500
ESSAIS = 2
MAX_CHARS_EVALUATION = 18_000   # évaluation relue (≈ 6 000 tokens) : marge confortable dans 16 k
MAX_CHARS_CONSIGNES = 1_500
MAX_PUCES = 6
TITRE_FEUILLE = "Fiches candidats"

# Valeurs qui ne disent rien de l'offre : à ne pas soumettre à l'IA comme des faits.
_VIDES = ("n/a", "", "—", "-")


def _est_exploitable(valeur: object) -> bool:
    texte = str(valeur or "").strip()
    return bool(texte) and texte.lower() not in _VIDES and not texte.startswith("⚠")


def _liste(valeur: object) -> list[str]:
    """Liste de puces propre : chaînes non vides, sans tiret de tête, plafonnée."""
    if isinstance(valeur, str):
        valeur = valeur.splitlines()
    if not isinstance(valeur, list):
        return []
    puces = []
    for v in valeur:
        texte = re.sub(r"^\s*[-*•]\s*", "", str(v or "")).replace("**", "").strip()
        if texte:
            puces.append(texte)
    return puces[:MAX_PUCES]


def lire_fiche(reponse: str) -> dict | None:
    """
    Lit `{"resume", "points_forts", "points_faibles", "vigilance"}`. Tolérant aux retours à la
    ligne bruts dans les chaînes (`strict=False`, cf. grille_analyse._lire_reponse). None si rien
    d'exploitable : ni résumé, ni point fort, ni point faible.
    """
    match = re.search(r"\{.*\}", reponse or "", re.DOTALL)
    data = None
    if match:
        try:
            data = json.loads(match.group(), strict=False)
        except json.JSONDecodeError:
            data = None
    if not isinstance(data, dict):
        # Pas de JSON lisible : le modèle a parfois rendu la fiche en Markdown (titres + puces).
        return _lire_fiche_texte(reponse or "")
    fiche = {
        "resume": str(data.get("resume") or "").replace("**", "").strip(),
        "points_forts": _liste(data.get("points_forts")),
        "points_faibles": _liste(data.get("points_faibles")),
        "vigilance": _liste(data.get("vigilance")),
    }
    if not (fiche["resume"] or fiche["points_forts"] or fiche["points_faibles"]):
        return None
    return fiche


_RUBRIQUES_TEXTE = (
    ("points_forts", ("points forts", "point fort", "forces")),
    ("points_faibles", ("points faibles", "point faible", "faiblesses", "axes d'amélioration")),
    ("vigilance", ("vigilance", "réserves", "à vérifier")),
    ("resume", ("résumé", "resume", "en bref", "synthèse")),
)


def _lire_fiche_texte(texte: str) -> dict | None:
    """Repli : fiche rendue en Markdown — une rubrique par titre (ou ligne « Points forts : »), puis ses puces."""
    fiche: dict = {"resume": "", "points_forts": [], "points_faibles": [], "vigilance": []}
    courante: str | None = None
    for brute in texte.splitlines():
        ligne = brute.strip()
        if not ligne:
            continue
        titre = re.sub(r"^[#*\s]+|[*:\s]+$", "", ligne).lower()
        rubrique = next((cle for cle, mots in _RUBRIQUES_TEXTE
                         if len(titre) < 40 and any(titre.startswith(m) or titre.endswith(m) for m in mots)), None)
        if rubrique:
            courante = rubrique
            continue
        if courante == "resume":
            fiche["resume"] = (fiche["resume"] + " " + ligne.replace("**", "")).strip()
        elif courante and re.match(r"^[-*•]\s+", ligne):
            fiche[courante] += _liste([ligne])
    for cle in ("points_forts", "points_faibles", "vigilance"):
        fiche[cle] = fiche[cle][:MAX_PUCES]
    if not (fiche["resume"] or fiche["points_forts"] or fiche["points_faibles"]):
        return None
    return fiche


def construire_prompt(nom: str, valeurs: dict[str, str], instructions: str | None) -> str | None:
    """Prompt de la fiche, ou None si l'évaluation ne contient rien d'exploitable."""
    lignes, reste = [], MAX_CHARS_EVALUATION
    for critere, valeur in valeurs.items():
        if not _est_exploitable(valeur):
            continue
        ligne = f"- {critere} : {str(valeur).strip()}"[:reste]
        lignes.append(ligne)
        reste -= len(ligne) + 1
        if reste <= 200:
            break
    if not lignes:
        return None
    evaluation = "\n".join(lignes)
    contexte = (f"\nContexte donné par l'analyste (pour le fond, pas pour le format) :\n"
                f"{instructions.strip()[:MAX_CHARS_CONSIGNES]}\n") if instructions and instructions.strip() else ""
    return f"""Tu es analyste des offres d'un marché public. Voici l'évaluation, critère par critère, de l'offre du candidat « {nom} » :

{evaluation}
{contexte}
=== CONSIGNE ===
Rédige la fiche de synthèse de CE candidat, en t'appuyant UNIQUEMENT sur l'évaluation ci-dessus (n'invente rien) :
- « resume » : 2 à 4 phrases factuelles sur l'offre ;
- « points_forts » : 2 à {MAX_PUCES} puces courtes, techniques ou financières ;
- « points_faibles » : 1 à {MAX_PUCES} puces courtes (manques, imprécisions, risques) ;
- « vigilance » : 0 à 3 puces (pièces manquantes, réserves, points à faire préciser).
Réponds UNIQUEMENT par un objet JSON, sans texte autour :
{{"resume": "…", "points_forts": ["…"], "points_faibles": ["…"], "vigilance": ["…"]}}"""


async def rediger_fiche(
    nom: str,
    valeurs: dict[str, str],
    instructions: str | None,
    model: str,
    ollama,
) -> dict | None:
    """Fiche du candidat `nom`, ou None (rien d'exploitable, appel ou réponse KO)."""
    prompt = construire_prompt(nom, valeurs, instructions)
    if prompt is None:
        return None
    for essai in range(1, ESSAIS + 1):
        try:
            reponse = await ollama.generate(prompt, model=model, num_predict=TOKENS_FICHE_MAX)
        except Exception as e:  # noqa: BLE001 — la fiche est un plus : son échec ne casse pas le tableau
            log.warning("Fiche candidat : appel IA impossible", essai=essai, erreur=str(e) or type(e).__name__)
            continue
        fiche = lire_fiche(reponse)
        if fiche is not None:
            return fiche
        # Diagnostic SANS contenu (offres confidentielles) : forme de la réponse seulement.
        texte = reponse or ""
        log.warning("Fiche candidat : réponse IA inexploitable", essai=essai, modele=model,
                    longueur=len(texte), json_ouvert="{" in texte, json_ferme=texte.rstrip().endswith(("}", "```")))
    return None


# ─── Rendus ──────────────────────────────────────────────────────────────────


def a_des_fiches(groupes: list[dict]) -> bool:
    """Vrai si la comparaison a demandé des fiches (les anciens résultats n'ont pas la clé)."""
    return any("fiche" in g for g in groupes)


def _sections(fiche: dict) -> list[tuple[str, list[str]]]:
    return [(titre, fiche.get(cle) or []) for titre, cle in (
        ("Points forts", "points_forts"), ("Points faibles", "points_faibles"),
        ("Points de vigilance", "vigilance"))]


def fiches_markdown(groupes: list[dict]) -> str:
    """Section « Fiches candidats » en Markdown (vide si aucune fiche demandée)."""
    if not a_des_fiches(groupes):
        return ""
    lignes = ["## Fiches candidats", ""]
    for g in groupes:
        lignes += [f"### {g.get('nom', '?')}", ""]
        fiche = g.get("fiche")
        if not fiche:
            lignes += ["_Fiche non produite par l'IA._", ""]
            continue
        if fiche.get("resume"):
            lignes += [fiche["resume"], ""]
        for titre, puces in _sections(fiche):
            if puces:
                lignes += [f"**{titre}**", ""] + [f"- {p}" for p in puces] + [""]
    return "\n".join(lignes).rstrip()


def ajouter_fiches_docx(doc, groupes: list[dict]) -> None:
    """Ajoute la section « Fiches candidats » à un document python-docx."""
    if not a_des_fiches(groupes):
        return
    doc.add_heading("Fiches candidats", level=1)
    for g in groupes:
        doc.add_heading(str(g.get("nom", "?")), level=2)
        fiche = g.get("fiche")
        if not fiche:
            doc.add_paragraph("Fiche non produite par l'IA.")
            continue
        if fiche.get("resume"):
            doc.add_paragraph(fiche["resume"])
        for titre, puces in _sections(fiche):
            if not puces:
                continue
            doc.add_paragraph().add_run(titre).bold = True
            for p in puces:
                doc.add_paragraph(p, style="List Bullet")


def ajouter_feuille_fiches(wb, groupes: list[dict]) -> None:
    """Ajoute une feuille « Fiches candidats » au classeur openpyxl (une colonne par candidat)."""
    if not a_des_fiches(groupes):
        return
    from openpyxl.styles import Alignment, Font

    titre = TITRE_FEUILLE if TITRE_FEUILLE not in wb.sheetnames else f"{TITRE_FEUILLE} (IA)"
    ws = wb.create_sheet(titre)
    ws.cell(row=1, column=1, value="Proposition IA — à valider par la commission").font = Font(italic=True)
    rubriques = [("Résumé", "resume"), ("Points forts", "points_forts"),
                 ("Points faibles", "points_faibles"), ("Points de vigilance", "vigilance")]
    for r, (libelle, _) in enumerate(rubriques, start=3):
        ws.cell(row=r, column=1, value=libelle).font = Font(bold=True)
    ws.column_dimensions["A"].width = 22
    for c, g in enumerate(groupes, start=2):
        ws.cell(row=2, column=c, value=str(g.get("nom", "?"))).font = Font(bold=True)
        ws.column_dimensions[ws.cell(row=2, column=c).column_letter].width = 60
        fiche = g.get("fiche")
        for r, (_, cle) in enumerate(rubriques, start=3):
            if not fiche:
                valeur = "Fiche non produite par l'IA." if r == 3 else ""
            elif cle == "resume":
                valeur = fiche.get("resume") or ""
            else:
                valeur = "\n".join(f"• {p}" for p in fiche.get(cle) or [])
            cellule = ws.cell(row=r, column=c, value=valeur)
            cellule.alignment = Alignment(vertical="top", wrap_text=True)
