"""
Grille d'analyse des offres — remplir le classeur Excel FOURNI, critère par critère.
=====================================================================================
Le tableau comparatif classique met un candidat par ligne et un critère par colonne. Une grille
d'analyse d'offres (RAO de marché public) fait l'inverse : les CRITÈRES sont en lignes (libellé,
documents à analyser, points), et chaque candidat occupe un BLOC de colonnes « Avis · Note sur
10 · Note finale » sous son nom (`{{NOM_SOCIETE}}`). La note finale y est souvent une formule.

Ce module :
  - reconnaît cette disposition (`detecter_grille`) — sans elle, le comparatif garde son
    comportement habituel ;
  - fait évaluer UN critère pour UN candidat par l'IA (`evaluer_critere`) : avis motivé + note ;
  - réécrit le classeur d'origine (`remplir_grille`) : nom, avis et note dans les cellules
    prévues, mise en forme et formules intactes.

Demande du 01/10/2026 : grille `MA26001_RAO_Lot1.xlsx`, offres CES et OTV.
"""
from __future__ import annotations

import json
import re
import unicodedata
from dataclasses import asdict, dataclass, field
from io import BytesIO
from pathlib import Path

from logger import get_logger

log = get_logger(__name__)

# Cellule écrite quand l'IA n'a rien rendu d'exploitable (même convention que le comparatif :
# un échec ne doit jamais passer pour un avis).
AVIS_ECHEC_IA = "⚠ Échec de l'IA"

LIGNES_ENTETE_MAX = 30          # l'en-tête « Avis / Note » est cherché dans les premières lignes

# Intitulés reconnus (normalisés par `_norm` : minuscules, sans accents). Le 01/10/2026, seule la
# grille MA26001 était connue (« Avis », « Note sur 10 », « Libellé ») ; le 11/10, la grille
# MA26041 n'a pas été reconnue et le comparatif l'a refusée (« aucune colonne en ligne 1 ») :
# d'une grille à l'autre, les mêmes colonnes portent d'autres mots.
_AVIS = ("avis", "commentaire", "appreciation", "observation", "justification", "motivation")
_NOTE = ("note", "notation", "score", "point")
_NOTE_EXCLUS = ("final", "total", "ponder", "maxi")
_LIBELLE = ("libelle", "critere", "sous critere", "sous-critere", "intitule", "designation")
_POIDS = ("point", "ponderation", "coef", "poids", "bareme")

# Nom de candidat GÉNÉRIQUE laissé par le modèle (« candidat #01 », « Entreprise 2 »…) : à remplacer
# par le vrai nom — avant, seuls les `{{…}}` et les cases vides l'étaient (grille MA26041, 11/10/2026).
_NOM_GENERIQUE = re.compile(r"^(candidat|societe|entreprise|soumissionnaire|offre|lot)s?\s*(n|no|#)?\s*\d*$")
_ECHELLE = re.compile(r"(?:sur|/)\s*(\d+(?:[.,]\d+)?)")
NOTE_MAX_DEFAUT = 10.0

# Barème écrit dans le classeur (30/10/2026 — ROADMAP « barème lu dans la grille ») : un niveau par
# cellule ou par ligne de cellule, « 0 = Absent », « 4 : Insuffisant », « 10 → Complet »… ; ou un
# petit tableau note | libellé sous un titre « Barème » / « Échelle de notation ». Sans lui, l'IA
# notait sur sa propre idée de « 7 sur 10 ».
_BAREME_NIVEAU = re.compile(
    r"^\s*(\d+(?:[.,]\d+)?)\s*(?:pts?|points?)?\s*(?:=|:|→|->|–|—)\s*(\S.{0,150}?)\s*$", re.I)
_BAREME_TITRE = ("bareme", "echelle", "grille de notation", "notation", "cotation")
BAREME_NIVEAUX_MAX = 12          # au-delà, ce n'est plus un barème mais un tableau de données
# Caractères de documents par évaluation. num_ctx = 16 384 tokens, ~3 caractères par token en
# français : 24 000 caractères ≈ 8 000 tokens, le reste pour les consignes et l'AVIS. À 40 000
# (01/10/2026), le prompt remplissait presque tout le contexte : plus de place pour répondre,
# fenêtre glissante, génération sans fin, 504 de la passerelle après 300 s — à chaque critère.
BUDGET_CRITERE = 24_000
# Part des « documents du marché » (CCTP, règlement) par critère : prise SUR le budget du critère,
# le prompt reste donc de la même taille (11/10/2026 — confronter l'offre aux exigences).
BUDGET_EXIGENCES = 6_000
MAX_CHARS_PAR_DOC = 12_000
TOKENS_AVIS_MAX = 1_000         # un avis de 3 à 6 phrases + la note : large

_PLACEHOLDER = re.compile(r"\{\{.*?\}\}")


def _norm(texte: object) -> str:
    """Minuscules, sans accents ni espaces superflus — pour comparer des en-têtes."""
    s = unicodedata.normalize("NFKD", str(texte or "")).encode("ascii", "ignore").decode()
    return " ".join(s.lower().split())


@dataclass
class Bloc:
    """Colonnes d'un candidat. `nom_cellule` = cellule du nom (souvent fusionnée au-dessus)."""
    nom_cellule: str
    col_avis: int
    col_note: int | None


@dataclass
class Critere:
    ligne: int
    libelle: str            # texte complet de la cellule (titre + sous-critères)
    documents: str          # « Documents analysés » — pièces à lire en priorité
    points: float | None    # « Points par sous critère »

    @property
    def titre(self) -> str:
        """Première ligne du libellé : sert de nom de critère (colonnes, Markdown, PDF)."""
        return (self.libelle.strip().splitlines() or [""])[0].strip()


@dataclass
class Grille:
    feuille: str
    ligne_entete: int
    blocs: list[Bloc] = field(default_factory=list)
    criteres: list[Critere] = field(default_factory=list)
    # Échelle lue dans l'en-tête (« Note sur 20 ») ; 10 à défaut. Avant le 11/10/2026, la note était
    # toujours bornée à 10 — sur une grille « sur 20 », l'IA ne pouvait pas dépasser la moitié.
    note_max: float = NOTE_MAX_DEFAUT
    # Niveaux du barème lus dans le classeur, triés : [{note, libelle}] ; [] si la grille n'en dit rien.
    bareme: list[dict] = field(default_factory=list)

    def to_dict(self) -> dict:
        return asdict(self)

    @classmethod
    def from_dict(cls, d: dict) -> "Grille":
        return cls(   # les clés en plus (`fichier` : copie du classeur) sont ignorées
            feuille=d["feuille"], ligne_entete=d["ligne_entete"],
            blocs=[Bloc(**b) for b in d.get("blocs", [])],
            criteres=[Critere(**c) for c in d.get("criteres", [])],
            note_max=float(d.get("note_max") or NOTE_MAX_DEFAUT),   # analyses lancées avant : sur 10
            bareme=list(d.get("bareme") or []),                    # avant le 30/10/2026 : pas de barème
        )


def _cellule_haut_gauche(ws, ligne: int, col: int) -> str:
    """Coordonnée de la cellule qui porte la valeur (la 1ʳᵉ d'une plage fusionnée)."""
    for plage in ws.merged_cells.ranges:
        if plage.min_row <= ligne <= plage.max_row and plage.min_col <= col <= plage.max_col:
            return ws.cell(row=plage.min_row, column=plage.min_col).coordinate
    return ws.cell(row=ligne, column=col).coordinate


def _est_note(entete: str) -> bool:
    """« Note sur 10 », « Notation », « Points obtenus »… mais ni la note finale ni le total."""
    return entete.startswith(_NOTE) and not any(x in entete for x in _NOTE_EXCLUS)


def _colonne_la_plus_textuelle(ws, ligne_entete: int, colonnes: list[int]) -> int | None:
    """Parmi `colonnes`, celle qui a le plus de cellules texte non vides sous l'en-tête (au moins 2)."""
    meilleure, max_textes = None, 1
    for c in colonnes:
        textes = 0
        for r in range(ligne_entete + 1, ws.max_row + 1):
            v = ws.cell(row=r, column=c).value
            textes += isinstance(v, str) and bool(v.strip())
        if textes > max_textes:
            meilleure, max_textes = c, textes
    return meilleure


def _echelle(entete: str) -> float | None:
    """« note sur 20 », « notation /20 » → 20.0 ; None si l'en-tête ne dit rien."""
    m = _ECHELLE.search(entete)
    return float(m.group(1).replace(",", ".")) if m else None


def _ancre_nom(ws, ligne: int, col: int, col_min: int) -> str:
    """
    Cellule qui porte le nom du candidat au-dessus de la colonne `col` : la plage fusionnée qui la
    couvre, sinon la dernière cellule non vide à sa gauche (jusqu'à `col_min`) — c'est ainsi que
    s'affiche un texte « centré sur plusieurs colonnes » (grille MA26041 : « candidat #01 » en C1
    couvre C à F, sans fusion). "" si rien.
    """
    fusion = _cellule_haut_gauche(ws, ligne, col)
    if ws[fusion].value not in (None, ""):
        return fusion
    for c in range(col, col_min - 1, -1):
        if ws.cell(row=ligne, column=c).value not in (None, ""):
            return ws.cell(row=ligne, column=c).coordinate
    return fusion


def _nombre_cellule(v: object) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str) and re.fullmatch(r"\s*\d+(?:[.,]\d+)?\s*", v):
        return float(v.replace(",", "."))
    return None


def lire_bareme(wb, exclure: set[tuple[str, str]] = frozenset()) -> list[dict]:
    """
    Niveaux du barème écrits n'importe où dans le classeur, triés par note : [{note, libelle}].
    Deux formes : « N = libellé » (une cellule, ou une ligne d'une cellule) ; tableau note | libellé
    sous un titre « Barème » / « Échelle »… `exclure` = cellules (feuille, coordonnée) déjà envoyées
    à l'IA autrement — les libellés de critères, qui peuvent contenir leur propre « 0 = … ».
    Rend [] s'il y a moins de 2 niveaux (un « 3 : … » isolé n'est pas un barème).
    """
    niveaux: dict[float, str] = {}
    for ws in wb.worksheets:
        for ligne in ws.iter_rows():
            for cell in ligne:
                v = cell.value
                if not isinstance(v, str) or (ws.title, cell.coordinate) in exclure:
                    continue
                for morceau in v.splitlines():
                    m = _BAREME_NIVEAU.match(morceau)
                    if m:
                        niveaux.setdefault(float(m.group(1).replace(",", ".")), m.group(2).strip())
                # Tableau sous un titre : nombre dans la colonne du titre, libellé à sa droite.
                titre = _norm(v)
                if len(titre) > 40 or not titre.startswith(_BAREME_TITRE):
                    continue
                trouves: dict[float, str] = {}
                for r in range(cell.row + 1, min(cell.row + BAREME_NIVEAUX_MAX + 1, ws.max_row) + 1):
                    note = _nombre_cellule(ws.cell(row=r, column=cell.column).value)
                    libelle = ws.cell(row=r, column=cell.column + 1).value
                    if note is None or not isinstance(libelle, str) or not libelle.strip() \
                            or len(libelle) > 150 or _nombre_cellule(libelle) is not None:
                        break
                    trouves.setdefault(note, libelle.strip())
                if len(trouves) >= 2:
                    for note, libelle in trouves.items():
                        niveaux.setdefault(note, libelle)
    if len(niveaux) < 2 or len(niveaux) > BAREME_NIVEAUX_MAX:
        return []
    return [{"note": n, "libelle": niveaux[n]} for n in sorted(niveaux)]


def _ouvrir(chemin: Path, data_only: bool = False):
    import openpyxl
    return openpyxl.load_workbook(str(chemin), data_only=data_only)


def detecter_grille(chemin: Path) -> Grille | None:
    """
    Reconnaît une grille d'analyse : une ligne d'en-tête avec au moins une cellule « Avis »
    suivie d'une « Note … » (hors « note finale »), et une colonne de libellés de critères.
    Rend None pour tout autre classeur — le comparatif classique s'applique alors.
    """
    try:
        wb = _ouvrir(chemin)
    except Exception as e:  # noqa: BLE001 — classeur illisible : pas une grille
        log.warning("Grille : classeur illisible", erreur=str(e))
        return None

    for ws in wb.worksheets:
        for ligne in range(1, min(ws.max_row, LIGNES_ENTETE_MAX) + 1):
            entetes = {c: _norm(ws.cell(row=ligne, column=c).value) for c in range(1, ws.max_column + 1)}
            cols_avis = [c for c, v in entetes.items() if v.startswith(_AVIS)]
            if not cols_avis:
                continue

            blocs: list[Bloc] = []
            for i, col_avis in enumerate(cols_avis):
                fin = cols_avis[i + 1] if i + 1 < len(cols_avis) else ws.max_column + 1
                col_note = next((c for c in range(col_avis + 1, fin) if _est_note(entetes[c])), None)
                nom = _ancre_nom(ws, ligne - 1, col_avis, cols_avis[0]) if ligne > 1 else ""
                # Plusieurs paires « Avis / Note » sous le MÊME nom de candidat : seule la 1re est à
                # l'IA, les suivantes restent à la commission (MA26041 : 10 paires, 5 candidats).
                if nom and any(b.nom_cellule == nom for b in blocs):
                    continue
                blocs.append(Bloc(nom_cellule=nom, col_avis=col_avis, col_note=col_note))
            if not any(b.col_note for b in blocs):
                continue
            col_note_1 = next(b.col_note for b in blocs if b.col_note)
            echelle_entete = _echelle(entetes[col_note_1])
            note_max = echelle_entete or NOTE_MAX_DEFAUT

            avant = [c for c in entetes if c < cols_avis[0]]
            col_libelle = next((c for c in avant if entetes[c].startswith(_LIBELLE)), None)
            col_docs = next((c for c in avant if "document" in entetes[c]), None)
            col_points = next((c for c in avant if entetes[c].startswith(_POIDS)), None)
            if col_libelle is None:
                # Aucun intitulé reconnu : les libellés sont dans la colonne qui, avant les blocs,
                # porte le plus de texte sous l'en-tête.
                col_libelle = _colonne_la_plus_textuelle(
                    ws, ligne, [c for c in avant if c not in (col_docs, col_points)])
            if col_libelle is None:
                continue

            criteres: list[Critere] = []
            for r in range(ligne + 1, ws.max_row + 1):
                libelle = ws.cell(row=r, column=col_libelle).value
                if not isinstance(libelle, str) or not libelle.strip():
                    continue
                points = ws.cell(row=r, column=col_points).value if col_points else None
                criteres.append(Critere(
                    ligne=r,
                    libelle=libelle.strip(),
                    documents=str(ws.cell(row=r, column=col_docs).value or "").strip() if col_docs else "",
                    points=float(points) if isinstance(points, (int, float)) else None,
                ))
            if criteres:
                exclure = {(ws.title, ws.cell(row=c.ligne, column=col_libelle).coordinate) for c in criteres}
                bareme = lire_bareme(wb, exclure)
                if bareme and echelle_entete is None:
                    # « Note » sans « sur N » : le barème donne l'échelle (s'il va au-delà de 10).
                    note_max = max(note_max, bareme[-1]["note"])
                # Un niveau au-delà de l'échelle de la grille la contredit : il n'est pas repris.
                bareme = [b for b in bareme if 0 <= b["note"] <= note_max]
                return Grille(feuille=ws.title, ligne_entete=ligne, blocs=blocs, criteres=criteres,
                              note_max=note_max, bareme=bareme if len(bareme) >= 2 else [])
    return None


def texte_bareme(bareme: list[dict]) -> str:
    """« - 0 = Absent\n- 10 = Complet » — tel qu'il est envoyé à l'IA et affiché."""
    return "\n".join(f"- {b['note']:g} = {b['libelle']}" for b in bareme)


@dataclass
class CritereListe:
    titre: str
    points: float | None

    @property
    def colonne(self) -> str:
        """Nom de colonne du comparatif : la pondération y figure, l'IA comme la commission la voient."""
        return f"{self.titre} ({self.points:g} pts)" if self.points is not None else self.titre


def lire_liste_criteres(chemin: Path) -> list[CritereListe]:
    """
    Grille SANS bloc « Avis / Note » par candidat : une simple liste de critères EN LIGNES, sous un
    en-tête « Critère » / « Libellé » / « Intitulé »… placé n'importe où dans les premières lignes,
    avec une éventuelle colonne de pondération. Rend [] si rien de tel (au moins 2 critères) ;
    sinon le comparatif classique s'applique, ces critères pour colonnes.
    """
    try:
        wb = _ouvrir(chemin, data_only=True)   # pondérations éventuellement calculées par formule
    except Exception as e:  # noqa: BLE001 — classeur illisible
        log.warning("Grille : classeur illisible", erreur=str(e))
        return []

    for ws in wb.worksheets:
        for ligne in range(1, min(ws.max_row, LIGNES_ENTETE_MAX) + 1):
            entetes = {c: _norm(ws.cell(row=ligne, column=c).value) for c in range(1, ws.max_column + 1)}
            col_libelle = next((c for c, v in entetes.items() if v.startswith(_LIBELLE)), None)
            # Un TITRE (« Critères d'attribution », seul sur sa ligne) n'est pas un en-tête.
            if col_libelle is None or sum(1 for v in entetes.values() if v) < 2:
                continue
            col_points = next((c for c, v in entetes.items()
                               if c != col_libelle and v.startswith(_POIDS)), None)
            criteres: list[CritereListe] = []
            for r in range(ligne + 1, ws.max_row + 1):
                valeur = ws.cell(row=r, column=col_libelle).value
                if not isinstance(valeur, str) or not valeur.strip():
                    continue
                titre = valeur.strip().splitlines()[0].strip()
                if _norm(titre).startswith(("total", "sous-total", "sous total", "note finale")):
                    continue
                points = ws.cell(row=r, column=col_points).value if col_points else None
                criteres.append(CritereListe(
                    titre=titre, points=float(points) if isinstance(points, (int, float)) else None))
            if len(criteres) >= 2:
                return criteres
    return []


def apercu_classeur(chemin: Path) -> dict:
    """
    Ce que le Tableau comparatif fera de ce classeur — affiché AVANT le lancement, pour qu'une
    grille mal reconnue se voie tout de suite (et pas par un refus au lancement, cf. 11/10/2026).
    `type` : grille (remplie telle quelle) · liste (critères en lignes, tableau à part) ·
    colonnes (en-têtes de la ligne 1) · inconnu.
    """
    grille = detecter_grille(chemin)
    if grille:
        bareme = (f" Barème lu : {len(grille.bareme)} niveaux, de {grille.bareme[0]['note']:g} "
                  f"({grille.bareme[0]['libelle']}) à {grille.bareme[-1]['note']:g} "
                  f"({grille.bareme[-1]['libelle']}) — l'IA notera selon ce barème."
                  if grille.bareme else " Aucun barème lu dans le classeur : précisez-le dans les consignes si besoin.")
        return {
            "type": "grille",
            "criteres": [{"nom": c.titre, "points": c.points, "documents": c.documents or None}
                         for c in grille.criteres],
            "nb_candidats": len(grille.blocs),
            "bareme": grille.bareme,
            "message": f"Grille d'analyse reconnue : {len(grille.criteres)} critère(s), "
                       f"{len(grille.blocs)} candidat(s) au plus, notes sur {grille.note_max:g}. "
                       "Elle sera rendue remplie, formules intactes." + bareme,
        }
    liste = lire_liste_criteres(chemin)
    if liste:
        return {
            "type": "liste",
            "criteres": [{"nom": c.titre, "points": c.points, "documents": None} for c in liste],
            "nb_candidats": None,
            "message": f"Liste de {len(liste)} critère(s) reconnue (en lignes). Le tableau sera produit "
                       "à part : la grille n'a pas de colonnes « Avis / Note » par candidat à remplir.",
        }
    entetes: list[str] = []
    try:
        ws = _ouvrir(chemin, data_only=True).active
        entetes = [str(v).strip() for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
                   if isinstance(v, str) and v.strip()]
    except Exception as e:  # noqa: BLE001 — classeur illisible
        log.warning("Grille : classeur illisible", erreur=str(e))
    if entetes:
        return {
            "type": "colonnes",
            "criteres": [{"nom": e, "points": None, "documents": None} for e in entetes],
            "nb_candidats": None,
            "message": f"{len(entetes)} colonne(s) lue(s) en ligne 1, utilisées comme critères.",
        }
    return {
        "type": "inconnu", "criteres": [], "nb_candidats": None,
        "message": "Classeur non reconnu : ni bloc « Avis / Note » par candidat, ni colonne « Critère » / "
                   "« Libellé », ni en-têtes en ligne 1. Choisissez « Je les saisis » pour taper les critères.",
    }


# ─── Évaluation par l'IA ─────────────────────────────────────────────────────


def _mots_cles(texte: str) -> set[str]:
    return {m for m in re.findall(r"[a-z0-9]+", _norm(texte)) if len(m) >= 4 or m.isdigit()}


def _ordonner_documents(docs: list, indications: str) -> list:
    """Les pièces citées dans « Documents analysés » d'abord (recoupement des mots du nom)."""
    cles = _mots_cles(indications)
    if not cles:
        return list(docs)
    return sorted(docs, key=lambda d: -len(cles & _mots_cles(d.nom)))


def _contexte(docs: list, budget: int) -> str:
    parts: list[str] = []
    reste = budget
    for doc in docs:
        texte = (doc.texte_extrait or "").strip()
        if not texte:
            continue
        texte = texte[:MAX_CHARS_PAR_DOC]
        entete = f"\n--- {doc.nom} ---\n"
        place = reste - len(entete)
        if place <= 200:
            break
        texte = texte[:place]
        parts.append(entete + texte)
        reste -= len(entete) + len(texte)
    return "\n".join(parts) if parts else "(aucun texte exploitable dans les documents de ce candidat)"


async def passages_pertinents(critere: Critere, docs: list, budget: int = BUDGET_CRITERE) -> str | None:
    """
    Les passages des documents du candidat les plus proches du critère (recherche sémantique
    sur les morceaux indexés, préfixe 1024-d). Un mémoire technique de 100 000 caractères ne
    tient pas dans le contexte : en prendre le début ratait les chapitres utiles. None si les
    vecteurs manquent (documents non vectorisés, Ollama indisponible) → repli sur `_contexte`.
    """
    from sqlalchemy import text

    from database import AsyncSessionLocal
    from routers.search import _embed_query
    from utils.vectors import matryoshka_prefix

    if not docs:
        return None
    requete = f"{critere.libelle}\n{critere.documents}"[:4000]
    vecteur = matryoshka_prefix(await _embed_query(requete))
    if vecteur is None:
        return None
    noms = {str(d.id): d.nom for d in docs}
    qs = "[" + ",".join(str(v) for v in vecteur) + "]"
    try:
        async with AsyncSessionLocal() as db:
            rows = (await db.execute(text("""
                SELECT CAST(document_id AS TEXT), chunk_index, chunk_text
                FROM embeddings
                WHERE embedding_small IS NOT NULL AND document_id = ANY(CAST(:ids AS uuid[]))
                ORDER BY embedding_small <=> CAST(:qs AS vector)
                LIMIT 120
            """), {"qs": qs, "ids": list(noms)})).fetchall()
    except Exception as e:  # noqa: BLE001 — base sans pgvector (tests) ou colonne absente : repli
        log.warning("Grille : recherche de passages impossible", erreur=str(e) or type(e).__name__)
        return None
    if not rows:
        return None

    retenus, reste = [], budget
    for did, idx, extrait in rows:
        extrait = (extrait or "").strip()
        if not extrait or len(extrait) + 80 > reste:
            continue
        retenus.append((did, idx, extrait))
        reste -= len(extrait) + 80
    # Relus dans l'ordre des documents : des passages d'un même mémoire se lisent mieux à la suite.
    retenus.sort(key=lambda r: (noms.get(r[0], ""), r[1]))
    parts, courant = [], None
    for did, idx, extrait in retenus:
        if did != courant:
            parts.append(f"\n--- {noms.get(did, '?')} (extraits) ---")
            courant = did
        parts.append(f"[…] {extrait}")
    return "\n".join(parts)


# Une note écrite « 15/20 », « 15 / 20 », « 15,5 sur 20 », « Note : 15 »… Le 11/10/2026, un prompt
# d'utilisateur demandait « la note sous la forme 15/20 » : 8 notes sur 14 étaient perdues, la
# conversion directe en nombre échouant sur « 15/20 ».
_NOTE_TEXTE = re.compile(r"(-?\d+(?:[.,]\d+)?)\s*(?:/|sur)\s*(\d+(?:[.,]\d+)?)|(-?\d+(?:[.,]\d+)?)", re.I)
# Note écrite en TÊTE d'avis (« 15/20 — Délais garantis… ») : récupérable quand le champ note manque.
_NOTE_EN_TETE = re.compile(r"^\s*(?:note\s*:?\s*)?(\d+(?:[.,]\d+)?)\s*(?:/|sur)\s*(\d+(?:[.,]\d+)?)", re.I)


def _nombre(texte: str) -> float:
    return float(texte.replace(",", "."))


def _note(valeur: object, note_max: float = NOTE_MAX_DEFAUT) -> float | None:
    """
    Note bornée à [0, note_max]. Accepte un nombre, ou un texte « 15/20 » / « 15 sur 20 » / « 15 » ;
    une note donnée sur une AUTRE échelle (« 8/10 » pour une grille sur 20) est convertie.
    """
    if valeur is None or isinstance(valeur, bool):
        return None
    if isinstance(valeur, (int, float)):
        brut = float(valeur)
    else:
        m = _NOTE_TEXTE.search(str(valeur))
        if not m:
            return None
        if m.group(1):
            brut, echelle = _nombre(m.group(1)), _nombre(m.group(2))
            if echelle > 0 and echelle != note_max:
                brut = brut * note_max / echelle
        else:
            brut = _nombre(m.group(3))
    return round(max(0.0, min(note_max, brut)), 2)


def _note_depuis_avis(avis: str, note_max: float = NOTE_MAX_DEFAUT) -> float | None:
    """Note écrite au début de l'avis (« 15/20 — … »), ou None."""
    m = _NOTE_EN_TETE.match(avis or "")
    return _note(f"{m.group(1)}/{m.group(2)}", note_max) if m else None


def _nettoyer_avis(avis: str) -> str:
    """Avis lisible dans une cellule Excel : sans gras Markdown, sans lignes vides en série."""
    avis = avis.replace("\\n", "\n").replace("**", "")   # « \n » littéraux d'un JSON mal échappé
    lignes = [ligne.strip() for ligne in avis.splitlines()]
    propre: list[str] = []
    for ligne in lignes:
        if ligne or (propre and propre[-1]):
            propre.append(ligne)
    return "\n".join(propre).strip()


def _lire_reponse(reponse: str, note_max: float = NOTE_MAX_DEFAUT) -> tuple[str, float | None] | None:
    """
    Lit `{"avis": …, "note": …}`. Tolérant : le 01/10/2026, ministral-3 a rendu 6 avis sur 6
    avec des retours à la ligne BRUTS dans la chaîne (JSON strict invalide) — tout était jeté
    alors que le contenu était bon. `strict=False` les accepte ; à défaut (guillemets non
    échappés, réponse tronquée), les deux champs sont relevés à la main.
    """
    texte = reponse or ""
    data = None
    match = re.search(r"\{.*\}", texte, re.DOTALL)
    if match:
        try:
            data = json.loads(match.group(), strict=False)
        except json.JSONDecodeError:
            data = None
    if isinstance(data, dict):
        cles = {str(k).strip().lower(): v for k, v in data.items()}
        avis = str(cles.get("avis") or cles.get("justification") or cles.get("commentaire") or "")
        note = _note(cles.get("note", cles.get("score")), note_max)
        if not avis and note is None:
            # Réponse rangée SOUS LE NOM DU CRITÈRE (« {"Délais…": "15/20 — …"} ») : un prompt
            # d'utilisateur écrit pour le comparatif classique y pousse le modèle (11/10/2026 : un
            # candidat entier en échec). On prend la première valeur texte comme avis.
            avis = next((str(v) for v in data.values() if isinstance(v, str) and v.strip()), "")
    else:
        m_note = re.search(r'"note"\s*:\s*"?(-?\d+(?:[.,]\d+)?)', texte)
        m_avis = re.search(r'"avis"\s*:\s*"(.*?)"\s*,\s*"note"', texte, re.DOTALL)
        if m_avis:
            avis = m_avis.group(1)
        else:   # réponse coupée avant la note : on garde l'avis jusqu'où il va
            m_avis = re.search(r'"avis"\s*:\s*"(.*)', texte, re.DOTALL)
            avis = re.sub(r'"?\s*\}?\s*(```)?\s*$', "", m_avis.group(1)) if m_avis else ""
        note = _note(m_note.group(1), note_max) if m_note else None
    avis = _nettoyer_avis(avis)
    if note is None:
        note = _note_depuis_avis(avis, note_max)
    if not avis and note is None:
        return None
    return avis, note


async def evaluer_critere(
    candidat: str,
    critere: Critere,
    docs: list,
    instructions: str | None,
    model: str,
    ollama,
    contexte: str | None = None,
    note_max: float = NOTE_MAX_DEFAUT,
    exigences: str | None = None,
    bareme: list[dict] | None = None,
) -> tuple[str, float | None, bool]:
    """
    Rend `(avis, note sur note_max, ok)`. Échec → `(AVIS_ECHEC_IA, None, False)`, jamais une note inventée.
    `contexte` = passages déjà choisis (`passages_pertinents`) ; à défaut, début des pièces citées.
    `exigences` = passages des documents du marché (CCTP, règlement) sur ce critère : l'avis dit
    alors, exigence par exigence, si l'offre y répond.
    """
    if not contexte:
        budget = BUDGET_CRITERE - (BUDGET_EXIGENCES if exigences else 0)
        contexte = _contexte(_ordonner_documents(docs, critere.documents), budget)
    pieces = f"\nPièces à examiner en priorité :\n{critere.documents}\n" if critere.documents else ""
    bloc_exigences = (f"\nExigences du marché sur ce critère (extraits du CCTP / règlement de consultation) :\n"
                      f"{exigences}\n") if exigences else ""
    consigne_exigences = (
        "- Confronte l'offre aux exigences du marché ci-dessus : pour chacune, dis si l'offre y répond, "
        "y répond en partie, ou n'en dit rien. Une exigence non couverte pèse sur la note.\n"
    ) if exigences else ""
    consignes_utilisateur = f"\nConsignes de l'analyste : {instructions}\n" if instructions else ""
    bloc_bareme = (f"\nBarème de notation de la grille :\n{texte_bareme(bareme)}\n") if bareme else ""
    consigne_note = (
        f"- « note » : nombre de 0 à {note_max:g}, attribué selon le barème de la grille rappelé ci-dessus "
        "(une note entre deux niveaux est permise si l'offre est entre les deux).\n"
    ) if bareme else (
        f"- « note » : nombre de 0 à {note_max:g} (décimales autorisées) reflétant la qualité de l'offre sur CE critère.\n"
    )
    rappel_bareme = f"\nRappel du barème :\n{texte_bareme(bareme)}\n" if bareme else ""
    # Les consignes sont répétées APRÈS les documents : en cas de dépassement de contexte, Ollama
    # coupe le DÉBUT du prompt — ce qui a déjà vidé 1 220 enrichissements (16/09/2026).
    prompt = f"""Tu es analyste des offres d'un marché public. Candidat évalué : {candidat}.

Critère à noter :
{critere.libelle}
{pieces}{bloc_exigences}{bloc_bareme}
Documents de l'offre du candidat :
{contexte}

=== CONSIGNE ===
Évalue l'offre de « {candidat} » sur le critère ci-dessus, en t'appuyant UNIQUEMENT sur ces documents.
{consigne_note.rstrip()}{rappel_bareme}
- « avis » : 6 phrases AU MAXIMUM (moins de 900 caractères), factuelles, en français — points forts,
  points faibles, éléments manquants, en citant les pièces. Si une pièce attendue est absente, dis-le.
  Pas de titres ni de listes à puces : un paragraphe, c'est le contenu d'une cellule Excel.
{consigne_exigences}{consignes_utilisateur}
Les consignes de l'analyste règlent le CONTENU de l'avis ; elles ne changent pas ce format :
« note » est un NOMBRE seul (ex. 15, jamais « 15/{note_max:g} »), « avis » est du texte.
Réponds UNIQUEMENT par un objet JSON, la note EN PREMIER, sans texte autour : {{"note": 7, "avis": "…"}}"""
    # La note d'abord : si la réponse est coupée (plafond de tokens), c'est l'avis qui s'arrête en
    # route, pas la note qui manque (01/10/2026 : avis de 3 300 caractères, note perdue).

    reponse = ""
    try:
        reponse = await ollama.generate(prompt, model=model, num_predict=TOKENS_AVIS_MAX)
        lu = _lire_reponse(reponse, note_max)
        if lu:
            return lu[0] or AVIS_ECHEC_IA, lu[1], True
    except Exception as e:  # noqa: BLE001 — appel KO : échec explicite ci-dessous
        # Journal SANS nom de candidat ni contenu d'offre (confidentiel, 11/10/2026).
        log.warning("Grille : évaluation IA impossible", critere=critere.titre, erreur=str(e) or type(e).__name__)
        return AVIS_ECHEC_IA, None, False
    log.warning("Grille : réponse IA inexploitable", critere=critere.titre, modele=model,
                longueur=len(reponse or ""), json_ouvert="{" in (reponse or ""))
    return AVIS_ECHEC_IA, None, False


async def exigences_marche(critere: Critere, docs_marche: list) -> str | None:
    """
    Passages des documents du marché (CCTP, règlement…) les plus proches du critère — recherche
    sémantique, repli sur le début des pièces si les vecteurs manquent. None sans document du marché.
    """
    if not docs_marche:
        return None
    passages = await passages_pertinents(critere, docs_marche, budget=BUDGET_EXIGENCES)
    if passages:
        return passages
    texte = _contexte(docs_marche, BUDGET_EXIGENCES)
    return None if texte.startswith("(aucun texte") else texte


# ─── Rendu ───────────────────────────────────────────────────────────────────


def remplir_grille(chemin: Path, grille: Grille, groupes: list[dict], classement: dict | None = None) -> bytes:
    """
    Écrit nom / avis / note de chaque candidat dans SON bloc (1ᵉʳ groupe → 1ᵉʳ bloc…).
    `groupes[i]` = {nom, evaluations: {str(ligne): {avis, note}}}. Les cellules à formule
    (note finale) ne sont jamais écrasées ; un nom déjà saisi est conservé, sauf s'il est générique
    (`{{…}}`, « candidat #01 »…). Une ligne « TOTAL … provisoire » sous le bloc reçoit la somme.
    """
    import openpyxl
    from openpyxl.styles import Alignment

    wb = openpyxl.load_workbook(str(chemin))
    ws = wb[grille.feuille]

    for bloc, groupe in zip(grille.blocs, groupes):
        if bloc.nom_cellule:
            cellule = ws[bloc.nom_cellule]
            actuel = cellule.value
            if actuel is None or (isinstance(actuel, str) and (
                    _PLACEHOLDER.search(actuel) or not actuel.strip() or _NOM_GENERIQUE.match(_norm(actuel)))):
                cellule.value = groupe.get("nom", "")
        _total_provisoire(ws, grille, bloc)

        evaluations = groupe.get("evaluations") or {}
        for critere in grille.criteres:
            ev = evaluations.get(str(critere.ligne))
            if not ev:
                continue
            c_avis = ws.cell(row=critere.ligne, column=bloc.col_avis)
            if not (isinstance(c_avis.value, str) and c_avis.value.startswith("=")):
                c_avis.value = ev.get("avis")
                c_avis.alignment = Alignment(vertical="top", wrap_text=True)
            if bloc.col_note and ev.get("note") is not None:
                c_note = ws.cell(row=critere.ligne, column=bloc.col_note)
                if not (isinstance(c_note.value, str) and c_note.value.startswith("=")):
                    c_note.value = ev["note"]

    # Classement (formules vers les cases de notes) et fiches : des feuilles EN PLUS, la grille
    # elle-même reste intacte.
    from services.classement import ajouter_feuille_classement
    from services.fiches_candidats import ajouter_feuille_fiches
    ajouter_feuille_classement(wb, classement, grille.to_dict(), groupes, grille.blocs[:len(groupes)])
    ajouter_feuille_fiches(wb, groupes)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _total_provisoire(ws, grille: Grille, bloc: Bloc) -> None:
    """
    Sous les critères, une ligne dont la colonne « Avis » du bloc dit « TOTAL … provisoire » (ou
    « TOTAL » seul) reçoit, dans la colonne « Note », la SOMME des notes du bloc — en formule, pour
    rester juste si la commission corrige une note. Un total « final / recalculé » est laissé à la
    commission ; une cellule déjà remplie n'est jamais écrasée. (Grille MA26041, 11/10/2026.)
    """
    if not bloc.col_note or not grille.criteres:
        return
    premiere = min(c.ligne for c in grille.criteres)
    derniere = max(c.ligne for c in grille.criteres)
    lettre = ws.cell(row=premiere, column=bloc.col_note).column_letter
    for r in range(derniere + 1, min(derniere + 6, ws.max_row) + 1):
        libelle = _norm(ws.cell(row=r, column=bloc.col_avis).value)
        if not libelle.startswith("total") or any(x in libelle for x in ("final", "recalcul")):
            continue
        cible = ws.cell(row=r, column=bloc.col_note)
        if cible.value in (None, ""):
            cible.value = f"=SUM({lettre}{premiere}:{lettre}{derniere})"
        return


def valeur_lisible(avis: str, note: float | None, note_max: float = NOTE_MAX_DEFAUT) -> str:
    """Cellule des rendus Markdown / PDF / Word : la note d'abord, puis l'avis."""
    if note is None:
        return avis
    return f"{note:g}/{note_max:g} — {avis}"
