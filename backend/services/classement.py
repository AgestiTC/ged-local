"""
Total pondéré et classement du Tableau comparatif (15/10/2026)
==============================================================
Calculés par le CODE, jamais par l'IA : chaque note est ramenée à ses points (note / échelle ×
points du critère), puis total, maximum possible, pourcentage et rang (ex aequo partagés).
Un candidat incomplet (note manquante, critère en échec) est classé avec son total partiel et
signalé — jamais présenté comme complet.

Calculé À LA LECTURE (rien n'est stocké) : vaut pour les comparaisons déjà faites et suit les
réparations. Dans l'Excel d'une grille, la feuille « Classement » porte de VRAIES formules vers les
cases de notes : une note corrigée par la commission recalcule total et rang.

Deux sources de notes :
  - grille d'analyse : `evaluations[ligne].note` sur `grille.note_max`, points = colonne de la grille ;
  - comparatif classique : note écrite en tête de valeur (« 15/20 — … »), points lus dans le nom
    de colonne (« Prix (40 pts) »).
"""
from __future__ import annotations

import re

_POINTS_COLONNE = re.compile(r"\((\d+(?:[.,]\d+)?)\s*pts?\)\s*$", re.I)
_NOTE_EN_TETE = re.compile(r"^\s*(?:note\s*:?\s*)?(\d+(?:[.,]\d+)?)\s*(?:/|sur)\s*(\d+(?:[.,]\d+)?)", re.I)
TITRE_FEUILLE = "Classement"


def _nombre(texte: str) -> float:
    return float(texte.replace(",", "."))


def _ratio_depuis_texte(valeur: object) -> float | None:
    """« 15/20 — … » → 0.75 ; None si la valeur ne commence pas par une note."""
    m = _NOTE_EN_TETE.match(str(valeur or ""))
    if not m or _nombre(m.group(2)) <= 0:
        return None
    return max(0.0, min(1.0, _nombre(m.group(1)) / _nombre(m.group(2))))


def criteres_et_ratios(groupes: list[dict], colonnes: list[str], grille: dict | None) -> tuple[list[dict], list[list[float | None]]]:
    """Critères `{titre, points, ligne}` et, par candidat, le ratio obtenu (0-1) de chaque critère (None = pas de note)."""
    if grille:
        note_max = float(grille.get("note_max") or 10) or 10.0
        criteres = [{"titre": (c.get("libelle") or "").strip().splitlines()[0] if c.get("libelle") else "",
                     "points": c.get("points"), "ligne": c.get("ligne")} for c in grille.get("criteres") or []]
        ratios = []
        for g in groupes:
            evs = g.get("evaluations") or {}
            ligne_ratios = []
            for c in criteres:
                note = (evs.get(str(c["ligne"])) or {}).get("note")
                ligne_ratios.append(None if note is None else max(0.0, min(1.0, float(note) / note_max)))
            ratios.append(ligne_ratios)
        return criteres, ratios

    criteres = []
    for col in colonnes:
        m = _POINTS_COLONNE.search(col)
        criteres.append({"titre": col, "points": _nombre(m.group(1)) if m else None, "ligne": None})
    ratios = [[_ratio_depuis_texte((g.get("valeurs") or {}).get(c["titre"])) for c in criteres] for g in groupes]
    return criteres, ratios


def calculer(groupes: list[dict], colonnes: list[str], grille: dict | None) -> dict | None:
    """
    `{pondere, max, lignes: [{nom, total, max, pct, rang, notes_manquantes, incomplet}]}` triées par rang,
    ou None si aucune note n'est exploitable (rien à classer : comparatif sans notation).
    """
    criteres, ratios = criteres_et_ratios(groupes, colonnes, grille)
    if not criteres or all(r is None for ligne in ratios for r in ligne):
        return None
    pondere = any(c["points"] for c in criteres)
    # Sans pondération : chaque critère pèse l'échelle de la note (sur 20 → 20 points chacun).
    echelle = float((grille or {}).get("note_max") or 20) or 20.0
    poids = [float(c["points"]) if pondere and c["points"] else (0.0 if pondere else echelle) for c in criteres]
    maximum = round(sum(poids), 2)

    lignes = []
    for g, ligne in zip(groupes, ratios):
        manquantes = sum(1 for r in ligne if r is None)
        total = round(sum((r or 0.0) * w for r, w in zip(ligne, poids)), 2)
        lignes.append({
            "nom": g.get("nom") or "?",
            "total": total, "max": maximum,
            "pct": round(100 * total / maximum, 1) if maximum else 0.0,
            "notes_manquantes": manquantes,
            "incomplet": bool(manquantes or g.get("echec_ia")),
        })
    # Rang « olympique » : ex aequo au même rang, le suivant saute (1, 1, 3).
    tries = sorted(lignes, key=lambda x: -x["total"])
    for i, ligne in enumerate(tries):
        ligne["rang"] = tries[i - 1]["rang"] if i and ligne["total"] == tries[i - 1]["total"] else i + 1
    return {"pondere": pondere, "max": maximum, "lignes": tries}


# ─── Rendus ──────────────────────────────────────────────────────────────────

_MEDAILLES = {1: "🥇", 2: "🥈", 3: "🥉"}


def classement_markdown(classement: dict | None) -> str:
    """Section « 🏆 Classement provisoire » (tableau Markdown : rendu par les exports PDF / Word)."""
    if not classement:
        return ""
    lignes = ["## 🏆 Classement provisoire", "",
              "| Rang | Candidat | Total | % |", "|---|---|---|---|"]
    for l in classement["lignes"]:
        rang = f"{_MEDAILLES.get(l['rang'], '')} {l['rang']}".strip()
        alerte = f" ⚠️ {l['notes_manquantes']} note(s) manquante(s)" if l["incomplet"] else ""
        lignes.append(f"| {rang} | {l['nom']}{alerte} | {l['total']:g} / {l['max']:g} | {l['pct']:g} % |")
    lignes += ["", "_Calculé par Matothèque à partir des notes proposées par l'IA — à valider par la commission._"]
    return "\n".join(lignes)


def ajouter_feuille_classement(wb, classement: dict | None, grille: dict | None = None,
                               groupes: list[dict] | None = None, blocs: list | None = None) -> None:
    """
    Feuille « Classement ». Avec une grille remplie (`grille` + `blocs`), le total de chaque candidat
    est une FORMULE vers ses cases de notes, et le rang une formule RANK : une note corrigée dans la
    grille recalcule tout. Sinon, valeurs calculées.
    """
    if not classement:
        return
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    titre = TITRE_FEUILLE if TITRE_FEUILLE not in wb.sheetnames else f"{TITRE_FEUILLE} (IA)"
    ws = wb.create_sheet(titre)
    ws.append(["Rang", "Candidat", "Total", "Sur", "%", "Notes manquantes"])
    for cellule in ws[1]:
        cellule.font = Font(bold=True)
    for col, largeur in zip("ABCDEF", (8, 34, 10, 8, 8, 18)):
        ws.column_dimensions[col].width = largeur

    par_nom = {l["nom"]: l for l in classement["lignes"]}
    noms = [g.get("nom") or "?" for g in groupes] if groupes else [l["nom"] for l in classement["lignes"]]
    premiere, derniere = 2, 1 + len(noms)
    feuille = (grille or {}).get("feuille")
    note_max = float((grille or {}).get("note_max") or 10) or 10.0
    criteres = (grille or {}).get("criteres") or []
    pondere = classement["pondere"]

    for i, nom in enumerate(noms):
        r = premiere + i
        l = par_nom.get(nom, {"total": 0, "notes_manquantes": 0})
        ws.cell(row=r, column=2, value=nom)
        bloc = blocs[i] if blocs and i < len(blocs) else None
        if feuille and bloc is not None and getattr(bloc, "col_note", None):
            lettre = get_column_letter(bloc.col_note)
            termes = [f"'{feuille}'!{lettre}{c['ligne']}" + (f"*{float(c['points']):g}" if pondere else "")
                      for c in criteres if not pondere or c.get("points")]
            somme = "+".join(termes) or "0"
            ws.cell(row=r, column=3, value=f"=ROUND(({somme})/{note_max:g},2)" if pondere else f"=ROUND({somme},2)")
        else:
            ws.cell(row=r, column=3, value=l["total"])
        ws.cell(row=r, column=4, value=classement["max"])
        ws.cell(row=r, column=5, value=f"=IF(D{r}>0,ROUND(100*C{r}/D{r},1),0)")
        ws.cell(row=r, column=1, value=f"=RANK(C{r},$C${premiere}:$C${derniere})")
        ws.cell(row=r, column=6, value=l["notes_manquantes"] or None)
    pied = ws.cell(row=derniere + 2, column=1,
                   value="Calculé à partir des notes proposées par l'IA — à valider par la commission. "
                         "Les totaux suivent les notes de la grille : corrigez une note, le classement se met à jour.")
    pied.font = Font(italic=True)
